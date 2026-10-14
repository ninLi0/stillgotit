import json
import pathlib
import sys
import zipfile
import fitz
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from PIL import Image, ImageDraw

root = pathlib.Path(sys.argv[1]); root.mkdir(parents=True, exist_ok=True)
content = "Matrix multiplication\nMultiply each row of A by each column of B. The inner dimensions must match."
(root / "lesson.md").write_text(content)
pdf = fitz.open(); page = pdf.new_page(); page.insert_text((60, 70), content, fontsize=13)
page.draw_rect(fitz.Rect(70, 150, 140, 210)); page.insert_text((78, 175), "1  2\n3  4")
pdf.save(root / "lesson.pdf"); pdf.close()
document = Document(); document.add_heading("Matrix multiplication", 0); document.add_paragraph(content); document.save(root / "lesson.docx")
presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text = "Matrix multiplication"; slide.placeholders[1].text = content; presentation.save(root / "lesson.pptx")
workbook = Workbook(); sheet = workbook.active; sheet.title = "Matrices"; sheet.append(["Matrix A", "Column 1", "Column 2"]); sheet.append(["Row 1", 1, 2]); sheet.append(["Row 2", 3, 4]); workbook.save(root / "lesson.xlsx")
(root / "lesson.ipynb").write_text(json.dumps({"cells": [{"cell_type": "markdown", "source": [content]}, {"cell_type": "code", "source": ["A = [[1, 2], [3, 4]]"]}]}))
image = Image.new("RGB", (480, 180), "white"); ImageDraw.Draw(image).text((20, 20), content, fill="black"); image.save(root / "lesson.png")
with zipfile.ZipFile(root / "lesson.zip", "w") as archive: archive.write(root / "lesson.md", "unit/lesson.md")
with zipfile.ZipFile(root / "unsafe.zip", "w") as archive: archive.writestr("../outside.txt", "Do not extract this.")
